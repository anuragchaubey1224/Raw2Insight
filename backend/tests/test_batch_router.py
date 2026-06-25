"""TestClient tests for the batch read/download endpoints (auth + db overridden)."""
import io
import json
import os
import tempfile
import uuid

os.environ["ENVIRONMENT"] = "development"
os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.NamedTemporaryFile(suffix='.db', delete=False).name}"

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import batch_router, batch_service, database
from app.auth import get_current_user
from app.database import Document, User, get_db


@pytest.fixture(scope="module", autouse=True)
def _init_db():
    database.init_db()


@pytest.fixture
def session():
    s = database.SessionLocal()
    yield s
    s.close()


@pytest.fixture
def user(session):
    u = User(email=f"{uuid.uuid4().hex[:8]}@t.com", password_hash="x")
    session.add(u)
    session.commit()
    session.refresh(u)
    return u


@pytest.fixture
def client(session, user):
    app = FastAPI()
    app.include_router(batch_router.router)
    app.dependency_overrides[get_db] = lambda: session
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app)


def _seed(session, user, statuses):
    batch = batch_service.create_batch(session, user.id, total_count=len(statuses))
    for i, st in enumerate(statuses):
        extracted = json.dumps({
            "vendor": f"Vendor{i}", "date": "2026-06-20", "total": 100 + i,
            "items": [{"description": "X", "qty": 1, "unit_price": 10, "line_total": 10}],
        }) if st == "completed" else None
        session.add(Document(
            user_id=user.id, job_id=uuid.uuid4().hex, filename=f"f{i}.jpg",
            status=st, batch_id=batch.batch_id, extracted_data=extracted,
        ))
    session.commit()
    return batch


def test_status_endpoint(client, session, user):
    batch = _seed(session, user, ["completed", "failed"])
    r = client.get(f"/api/v1/batch/{batch.batch_id}/status")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2 and body["batch_id"] == batch.batch_id


def test_results_endpoint(client, session, user):
    batch = _seed(session, user, ["completed", "completed"])
    r = client.get(f"/api/v1/batch/{batch.batch_id}/results")
    assert r.status_code == 200
    rows = r.json()["rows"]
    assert len(rows) == 2
    assert rows[0]["vendor"] == "Vendor0"
    assert rows[0]["items"][0]["description"] == "X"


def test_download_csv(client, session, user):
    batch = _seed(session, user, ["completed"])
    r = client.get(f"/api/v1/batch/{batch.batch_id}/download?format=csv")
    assert r.status_code == 200
    assert r.text.startswith("﻿File,Vendor,Date,Total,Items,Status") or "File,Vendor" in r.text


def test_download_xlsx(client, session, user):
    batch = _seed(session, user, ["completed", "failed"])
    r = client.get(f"/api/v1/batch/{batch.batch_id}/download?format=xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Line Items"]


def test_bad_format_rejected(client, session, user):
    batch = _seed(session, user, ["completed"])
    assert client.get(f"/api/v1/batch/{batch.batch_id}/download?format=pdf").status_code == 422


def test_unknown_batch_404(client):
    assert client.get("/api/v1/batch/does-not-exist/status").status_code == 404


def test_other_users_batch_is_404(client, session):
    other = User(email=f"{uuid.uuid4().hex[:8]}@t.com", password_hash="x")
    session.add(other)
    session.commit()
    session.refresh(other)
    batch = batch_service.create_batch(session, other.id, total_count=1)
    # current client user != owner -> must not leak another user's batch
    assert client.get(f"/api/v1/batch/{batch.batch_id}/status").status_code == 404


# ---- upload endpoint (worker stubbed so no real YOLO/OCR runs) ------------------------------

def _img(name="a.jpg"):
    return ("files", (name, b"\xff\xd8\xff\xf0fakejpegbytes", "image/jpeg"))


def test_batch_upload_creates_batch(client, session, user, monkeypatch):
    monkeypatch.setattr(batch_router, "_run_batch", lambda *a, **k: None)  # don't run extraction
    r = client.post("/api/v1/batch/upload",
                    files=[_img("a.jpg"), _img("b.jpg")], data={"engine": "local"})
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2 and body["status"] == "processing" and "batch_id" in body
    # a Batch row + one Document per bill were created
    batch = batch_service.get_batch(session, body["batch_id"])
    assert batch.total_count == 2
    docs = batch_service.get_batch_documents(session, body["batch_id"])
    assert len(docs) == 2 and all(d.batch_id == body["batch_id"] for d in docs)


def test_batch_upload_rejects_unsupported_type(client, monkeypatch):
    monkeypatch.setattr(batch_router, "_run_batch", lambda *a, **k: None)
    r = client.post("/api/v1/batch/upload",
                    files=[("files", ("note.txt", b"hello", "text/plain"))],
                    data={"engine": "local"})
    assert r.status_code == 400


def test_batch_upload_enforces_cap(client, monkeypatch):
    monkeypatch.setattr(batch_router, "_run_batch", lambda *a, **k: None)
    files = [_img(f"{i}.jpg") for i in range(11)]  # 11 > MAX_BATCH_SIZE (10)
    r = client.post("/api/v1/batch/upload", files=files, data={"engine": "local"})
    assert r.status_code == 400
    assert "10" in r.json()["detail"]
